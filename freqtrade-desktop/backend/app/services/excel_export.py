"""Excel per-trade report export.

Sheets are derived from 回测指标需求.txt (one sheet per top-level section),
plus an overview sheet and an outlier list. Numeric columns get sample Z-scores
with soft/hard outlier highlighting.
"""

from __future__ import annotations

import math
import re
import statistics
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..config import APP_DIR


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SOFT_FILL = PatternFill("solid", fgColor="FFF3CD")
HARD_FILL = PatternFill("solid", fgColor="FFC7CE")
THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def parse_catalog(path: Path | None = None) -> list[dict[str, Any]]:
    path = path or (APP_DIR / "回测指标需求.txt")
    sections: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None
    if not path.exists():
        return sections
    title_re = re.compile(r"^[一二三四五六七八九十]+、")
    for raw in path.read_text(encoding="utf-8").splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        if title_re.match(line):
            current = {"title": line.strip(), "fields": [], "desc": {}}
            sections.append(current)
            continue
        if current is None or "\t" not in line:
            continue
        parts = [p.strip() for p in line.split("\t")]
        field = parts[0]
        if not field or field == "指标字段":
            continue
        desc = parts[1] if len(parts) > 1 else ""
        for name in [x.strip() for x in field.split("/") if x.strip()]:
            current["fields"].append(name)
            current["desc"][name] = desc
    return sections


def _num(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    try:
        v = float(value)
        return v if v == v else None
    except (TypeError, ValueError):
        return None


def _is_pct(field: str) -> bool:
    return field.endswith("_pct") or field in (
        "return_pct", "mfe_pct", "mae_pct", "giveback_pct",
        "recovery_from_mae_pct", "max_runup_pct", "max_drawdown_pct",
    )


def _z_scores(values: list[float]) -> tuple[float, float, list[float | None]]:
    if len(values) < 2:
        return 0.0, 0.0, [None] * len(values)
    mean = statistics.mean(values)
    std = statistics.stdev(values)
    if std == 0:
        return mean, std, [None] * len(values)
    return mean, std, [(v - mean) / std for v in values]


def build_report(
    rows: list[dict[str, Any]],
    sections: list[dict[str, Any]],
    run_params: dict[str, Any],
    z_soft: float = 2.0,
    z_hard: float = 3.0,
    notes: dict[str, str] | None = None,
) -> Workbook:
    notes = notes or {}
    clean_rows = [
        {k: (None if v is not None and str(v) == "<NA>" else v) for k, v in row.items()}
        for row in rows
    ]
    wb = Workbook()
    wb.remove(wb.active)

    # ---------- 总览 ----------
    ws = wb.create_sheet("总览")
    ws.append(["参数", "值"])
    for key, value in run_params.items():
        ws.append([key, value])
    ws.append(["导出时间", datetime.now(timezone.utc).isoformat(timespec="seconds")])
    numeric_fields: list[str] = []
    for section in sections:
        for field in section["fields"]:
            if field not in numeric_fields and any(
                _num(row.get(field)) is not None for row in clean_rows
            ):
                numeric_fields.append(field)
    ws.append([])
    ws.append(["数值指标", "样本数", "均值", "中位数", "标准差", "最小", "最大", "Q25", "Q75", f"离群(|Z|>={z_soft})"])
    outlier_summary: dict[str, int] = {}
    for field in numeric_fields:
        values = [_num(r.get(field)) for r in clean_rows]
        values = [v for v in values if v is not None]
        if len(values) < 2:
            continue
        _, _, zs = _z_scores(values)
        outliers = sum(1 for z in zs if z is not None and abs(z) >= z_soft)
        outlier_summary[field] = outliers
        qs = sorted(values)
        def q(p: float) -> float:
            idx = int(p * (len(qs) - 1))
            return qs[idx]
        ws.append([
            field, len(values), round(statistics.mean(values), 6),
            round(statistics.median(values), 6), round(statistics.stdev(values), 6),
            round(min(values), 6), round(max(values), 6),
            round(q(0.25), 6), round(q(0.75), 6), outliers,
        ])
    ws.freeze_panes = "A3"
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    # ---------- 分类 sheet ----------
    for section in sections:
        ws = wb.create_sheet(section["title"][:31])
        fields = section["fields"]
        ws.append(fields)
        ws.append([notes.get(f, "口径见需求文档/分析说明") for f in fields])
        ws.freeze_panes = "A3"
        for col_idx, field in enumerate(fields, start=1):
            header_cell = ws.cell(row=1, column=col_idx)
            header_cell.fill = HEADER_FILL
            header_cell.font = HEADER_FONT
            ws.cell(row=2, column=col_idx).font = Font(size=9, italic=True, color="666666")
            ws.column_dimensions[get_column_letter(col_idx)].width = max(14, min(38, (len(field) + 4)))
        numeric_stats: dict[int, tuple[float, float, int]] = {}
        for col_idx, field in enumerate(fields, start=1):
            values = [_num(r.get(field)) for r in clean_rows]
            values = [v for v in values if v is not None]
            if len(values) >= 2:
                mean = statistics.mean(values)
                std = statistics.stdev(values)
                outliers = sum(1 for v in values if std and abs((v - mean) / std) >= z_soft)
                numeric_stats[col_idx] = (mean, std, outliers)
        for row_idx, record in enumerate(clean_rows, start=3):
            for col_idx, field in enumerate(fields, start=1):
                value = record.get(field)
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.border = BORDER
                if isinstance(value, bool):
                    continue
                num = _num(value)
                if num is not None and col_idx in numeric_stats:
                    cell.number_format = "0.00%" if _is_pct(field) else "0.0000"
                    mean, std, _ = numeric_stats[col_idx]
                    z = (num - mean) / std if std else None
                    if z is not None and abs(z) >= z_hard:
                        cell.fill = HARD_FILL
                    elif z is not None and abs(z) >= z_soft:
                        cell.fill = SOFT_FILL
                elif value is None:
                    cell.value = "NA"
                    cell.font = Font(color="999999")
                ws.row_dimensions[row_idx].height = 16
        ws.auto_filter.ref = f"A1:{get_column_letter(len(fields))}{max(len(clean_rows) + 2, 2)}"
        ws.alignment = Alignment(vertical="top")

    # ---------- 附加因子（资金费/基差等） ----------
    catalog_fields = {f for section in sections for f in section["fields"]}
    if clean_rows:
        extras = [k for k in clean_rows[0] if k not in catalog_fields and not k.startswith("_")]
        if extras:
            ws = wb.create_sheet("资金费与基差（合约）")
            ws.append(extras)
            ws.append([notes.get(f, "附加计算字段") for f in extras])
            ws.freeze_panes = "A3"
            for col_idx, field in enumerate(extras, start=1):
                cell = ws.cell(row=1, column=col_idx)
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
            for row_idx, record in enumerate(clean_rows, start=3):
                for col_idx, field in enumerate(extras, start=1):
                    cell = ws.cell(row=row_idx, column=col_idx, value=record.get(field))
                    cell.border = BORDER

    # ---------- 离群清单 ----------
    ws = wb.create_sheet("离群清单")
    ws.append(["trade_id", "指标", "数值", "Z值", "阈值"])
    for header_cell in ws[1]:
        header_cell.fill = HEADER_FILL
        header_cell.font = HEADER_FONT
    for field, count in outlier_summary.items():
        values = [_num(r.get(field)) for r in clean_rows]
        values = [v for v in values if v is not None]
        _, _, zs = _z_scores(values)
        for record, z in zip(clean_rows, zs):
            if z is not None and abs(z) >= z_soft:
                ws.append([
                    record.get("trade_id", ""), field,
                    record.get(field), round(z, 3), f"soft{z_soft}/hard{z_hard}",
                ])
    ws.freeze_panes = "A2"
    return wb


def default_filename(strategy: str) -> str:
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%S")
    safe = re.sub(r"[^\w.-]", "_", strategy)
    return f"回测日志-{safe}-{stamp}.xlsx"
